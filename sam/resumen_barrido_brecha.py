"""Resume barrido_fino_brecha_SAM.csv y finalistas_brecha*.csv para el grupo:

  - casos_seleccionados_brecha_SAM.csv : casos representativos de cada palanca (todos los indicadores)
  - barrido_fino_brecha_SAM.xlsx       : el mismo resumen en Excel (sin el problema del punto decimal del CSV
                                         en Excel en español) + matrices de horas con déficit + barrido completo

Inversión adicional (ΔCAPEX): solo para COMPARAR casos, con los costos por defecto de SAM 2026.7.3
(Fuel Cell - Commercial, escala comercial de 200 kW, no presupuesto real):
batería 184 $/kWh + 667 $/kW · SOFC 10.000 $/kW · FV 1,63 $/Wdc
(módulo 0,43 + BOS 0,35 + mano de obra 0,18 + margen 0,25 + inversor 0,05/1,25 + ingeniería 0,30 + permisos 0,03 + red 0,05).
"""
import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

df = pd.read_csv("barrido_fino_brecha_SAM.csv", encoding="utf-8-sig")
df["dCAPEX_MUSD"] = ((df.bat_MWh - 160) * 1000 * 184 + (df.bat_MW - 20) * 1000 * 667
                     + (df.sofc_MW - 110) * 1000 * 10_000 + (df.pv_MWdc - 300) * 1e6 * 1.63) / 1e6
df["dCAPEX_MUSD"] = df["dCAPEX_MUSD"].round(0)

# (nombre, FV, SOFC, rampa, bat MW, bat MWh, carga desde SOFC)
CASOS = [
    ("Diseño entregado", 300, 110, 20, 20, 160, 0), ("A · bat 60/480", 300, 110, 20, 60, 480, 0), ("B · bat 60/640", 300, 110, 20, 60, 640, 1),
    ("bat 60/720", 300, 110, 20, 60, 720, 1), ("F1 · bat 70/640", 300, 110, 20, 70, 640, 1), ("bat 80/640", 300, 110, 20, 80, 640, 1),
    ("bat 70/720", 300, 110, 20, 70, 720, 1), ("bat 80/800", 300, 110, 20, 80, 800, 1),
    ("SOFC 115 · bat 60/560", 300, 115, 20, 60, 560, 1), ("SOFC 115 · bat 70/560", 300, 115, 20, 70, 560, 1), ("SOFC 115 · bat 70/640", 300, 115, 20, 70, 640, 1),
    ("G · SOFC 120 · bat 40/320", 300, 120, 20, 40, 320, 1), ("SOFC 120 · bat 60/400", 300, 120, 20, 60, 400, 1), ("SOFC 120 · bat 60/480", 300, 120, 20, 60, 480, 1),
    ("F2 · SOFC 120 · bat 70/480", 300, 120, 20, 70, 480, 1), ("SOFC 120 · bat 80/560", 300, 120, 20, 80, 560, 1),
    ("FV 350 · bat 70/640", 350, 110, 20, 70, 640, 1), ("FV 350 · bat 80/640", 350, 110, 20, 80, 640, 1), ("C · FV 400 · bat 60/480", 400, 110, 20, 60, 480, 1),
    ("FV 400 · bat 70/640", 400, 110, 20, 70, 640, 1), ("F3 = D · FV 400 · bat 80/640", 400, 110, 20, 80, 640, 1), ("FV 400 · bat 90/800", 400, 110, 20, 90, 800, 1),
    ("B · rampa 30", 300, 110, 30, 60, 640, 1), ("B · rampa 40", 300, 110, 40, 60, 640, 1),
    ("F1 · rampa 30", 300, 110, 30, 70, 640, 1), ("F1 · rampa 40", 300, 110, 40, 70, 640, 1),
    ("F2 · rampa 30", 300, 120, 30, 70, 480, 1), ("F2 · rampa 40", 300, 120, 40, 70, 480, 1),
    ("F3 · rampa 30", 400, 110, 30, 80, 640, 1), ("F3 · rampa 40", 400, 110, 40, 80, 640, 1),
    ("SOFC 120 · bat 40/320 · rampa 40", 300, 120, 40, 40, 320, 1),
    ("F1 sin carga desde SOFC", 300, 110, 20, 70, 640, 0), ("F2 sin carga desde SOFC", 300, 120, 20, 70, 480, 0),
    ("F3 sin carga desde SOFC", 400, 110, 20, 80, 640, 0),
]
COLS = ["caso", "pv_MWdc", "sofc_MW", "rampa_pct_h", "bat_MW", "bat_MWh", "carga_desde_sofc", "cobertura_pct", "deficit_GWh",
        "horas_deficit", "pct_horas_deficit", "deficit_max_MW", "racha_max_h", "dias_con_deficit", "excedente_GWh", "H2_t",
        "firme_MW", "terreno_ha", "ciclos_bat", "deficit_peor_mes_GWh", "dCAPEX_MUSD"]
k = ["pv_MWdc", "sofc_MW", "rampa_pct_h", "bat_MW", "bat_MWh", "carga_desde_sofc"]
filas = []
for nombre, *v in CASOS:
    r = df[(df[k] == v).all(axis=1)]
    assert len(r) == 1, nombre
    filas.append({**r.iloc[0].to_dict(), "caso": nombre})
sel = pd.DataFrame(filas)[COLS]
sel.to_csv("casos_seleccionados_brecha_SAM.csv", index=False, encoding="utf-8-sig")

fin = pd.read_csv("finalistas_brecha_SAM.csv", encoding="utf-8-sig")
mens = pd.read_csv("finalistas_brecha_mensual_SAM.csv", encoding="utf-8-sig")
with pd.ExcelWriter("barrido_fino_brecha_SAM.xlsx", engine="openpyxl") as xw:
    sel.to_excel(xw, sheet_name="Casos seleccionados", index=False)
    fin.to_excel(xw, sheet_name="Finalistas (control SAM)", index=False)
    mens.pivot(index="mes", columns="caso", values="horas_deficit").to_excel(xw, sheet_name="Mensual horas")
    mens.pivot(index="mes", columns="caso", values="deficit_GWh").to_excel(xw, sheet_name="Mensual GWh")
    # matrices de horas con déficit (filas: MW, columnas: MWh), con carga desde SOFC
    fila0, ws_name = 0, "Matrices horas déficit"
    for pv, fc, ra in [(300, 110, 20), (300, 110, 30), (300, 110, 40), (350, 110, 20), (400, 110, 20),
                       (300, 115, 20), (300, 120, 20), (300, 120, 30), (300, 120, 40)]:
        t = df[(df.pv_MWdc == pv) & (df.sofc_MW == fc) & (df.rampa_pct_h == ra) & (df.carga_desde_sofc == 1)]
        t = t.pivot(index="bat_MW", columns="bat_MWh", values="horas_deficit")
        pd.DataFrame([[f"FV {pv} MWdc · SOFC {fc} MW · rampa {ra} %/h · carga desde SOFC (filas: batería MW; columnas: MWh)"]]).to_excel(
            xw, sheet_name=ws_name, startrow=fila0, index=False, header=False)
        t.to_excel(xw, sheet_name=ws_name, startrow=fila0 + 1)
        fila0 += len(t) + 4
    df.to_excel(xw, sheet_name="Barrido completo", index=False)

    wb = xw.book
    regla = lambda: ColorScaleRule(start_type="min", start_color="63BE7B", mid_type="percentile", mid_value=50,
                                   mid_color="FFEB84", end_type="max", end_color="F8696B")
    ws = wb[ws_name]
    for r0 in range(1, fila0, 10):
        ws.conditional_formatting.add(f"B{r0 + 2}:H{r0 + 7}", regla())
    ws.column_dimensions["A"].width = 10
    for nombre in ["Casos seleccionados", "Finalistas (control SAM)", "Barrido completo"]:
        w = wb[nombre]; w.freeze_panes = "B2"
        for i, c in enumerate(w[1], 1):
            w.column_dimensions[get_column_letter(i)].width = max(10, min(34, len(str(c.value)) + 2))
    h = COLS.index("horas_deficit") + 1
    wb["Casos seleccionados"].conditional_formatting.add(
        f"{get_column_letter(h)}2:{get_column_letter(h)}{len(sel) + 1}", regla())
print(sel.to_string(index=False))
